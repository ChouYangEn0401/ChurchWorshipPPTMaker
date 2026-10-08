"""把整理好的歌單（store.resolve_service 的結果）做成 .pptx。

為了「之後換圖片很方便」，背景不是貼在每一張投影片上，而是放在投影片母片的「版面配置」裡：
每一個用到的背景做成一個版面配置，名字叫「背景：xxx」，同一首歌的每一頁都套它。
在 PowerPoint 裡：檢視 → 投影片母片 → 點「背景：xxx」→ 背景格式 → 換圖片，
用到這個背景的每一頁就一起換好了（暗色遮罩也在版面配置裡，會跟著保留）。

文字是一般文字方塊，字型同時設定西文和東亞字型（不然中文會掉回 PowerPoint 預設字）。
每一頁的備忘稿寫「歌名｜段落」，用簡報者檢視時比較好跟。
"""

from __future__ import annotations

import copy
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.slide import SlideLayoutPart
from pptx.util import Emu, Inches, Pt

from . import images as imagestore
from .lyrics import header_scale, lyrics_layout, title_page_size

SLIDE_SIZES = {"4:3": (10.0, 7.5), "16:9": (13.333, 7.5)}
ALIGN = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}
BLANK_LAYOUT_INDEX = 6

NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def _rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#FFFFFF").lstrip("#")
    if len(h) != 6:
        h = "FFFFFF"
    return RGBColor.from_string(h.upper())


# ---------------------------------------------------------------------------
# 版面配置（一個背景一個）
# ---------------------------------------------------------------------------

def _new_layout(prs, name: str):
    master = prs.slide_master
    base = prs.slide_layouts[BLANK_LAYOUT_INDEX]
    package = master.part.package
    partname = package.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
    element = copy.deepcopy(base._element)
    element.set("type", "cust")
    element.set("preserve", "1")
    csld = element.find(qn("p:cSld"))
    csld.set("name", name)
    sp_tree = csld.find(qn("p:spTree"))
    for shape in list(sp_tree):
        if shape.tag in (qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp")):
            sp_tree.remove(shape)
    part = SlideLayoutPart(partname, base.part.content_type, package, element)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    id_list = master._element.get_or_add_sldLayoutIdLst()
    used = [int(x.get("id")) for x in id_list.findall(qn("p:sldLayoutId"))]
    used += [int(x.get("id")) for x in prs.part._element.iter(qn("p:sldMasterId"))]
    entry = etree.SubElement(id_list, qn("p:sldLayoutId"))
    entry.set("id", str(max(used) + 1))
    entry.set(qn("r:id"), rid)
    return part.slide_layout


def _cover_crop(img_w: int, img_h: int, slide_w: float, slide_h: float) -> dict[str, int]:
    """像 CSS 的 background-size: cover：填滿畫面、不變形，多出來的從兩邊平均裁掉。單位是 1/1000 %。"""
    if not img_w or not img_h:
        return {}
    img_ratio, slide_ratio = img_w / img_h, slide_w / slide_h
    if abs(img_ratio - slide_ratio) < 0.005:
        return {}
    if img_ratio > slide_ratio:
        cut = (1 - slide_ratio / img_ratio) / 2
        return {"l": round(cut * 100000), "r": round(cut * 100000)}
    cut = (1 - img_ratio / slide_ratio) / 2
    return {"t": round(cut * 100000), "b": round(cut * 100000)}


def _set_layout_background(layout, theme: dict, image_file: Path | None, slide_w_in: float, slide_h_in: float):
    csld = layout._element.find(qn("p:cSld"))
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)
    if image_file is not None:
        _, rid = layout.part.get_or_add_image_part(str(image_file))
        crop = _cover_crop(theme["image"].get("width") or 0, theme["image"].get("height") or 0, slide_w_in, slide_h_in)
        src_rect = "<a:srcRect " + " ".join(f'{k}="{v}"' for k, v in crop.items()) + "/>"
        fill = (f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/>{src_rect}'
                f'<a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    else:
        fill = f'<a:solidFill><a:srgbClr val="{(theme.get("bgColor") or "#000000").lstrip("#").upper()}"/></a:solidFill>'
    bg = etree.fromstring(f'<p:bg {NS}><p:bgPr>{fill}<a:effectLst/></p:bgPr></p:bg>')
    csld.insert(0, bg)


def _add_layout_overlay(layout, width: int, height: int, alpha: float) -> None:
    """暗色遮罩：半透明黑色方塊放在版面配置上（換背景圖時會保留）。"""
    if not alpha or alpha <= 0:
        return
    sp_tree = layout._element.find(qn("p:cSld")).find(qn("p:spTree"))
    ids = [int(x.get("id")) for x in sp_tree.iter(qn("p:cNvPr"))] or [1]
    val = int(max(0.0, min(1.0, float(alpha))) * 100000)
    shape = etree.fromstring(
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{max(ids) + 1}" name="暗色遮罩"/><p:cNvSpPr/><p:nvPr userDrawn="1"/>'
        f'</p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="000000">'
        f'<a:alpha val="{val}"/></a:srgbClr></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
    sp_tree.append(shape)


# ---------------------------------------------------------------------------
# 文字
# ---------------------------------------------------------------------------

def _style_run(run, part: dict, size: float | None = None) -> None:
    font = run.font
    font.size = Pt(float(size or part["size"]))
    font.bold = bool(part.get("bold"))
    font.color.rgb = _rgb(part.get("color"))
    font.name = part.get("font") or "Microsoft JhengHei"  # 建立 <a:latin>
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:effectLst", "a:ea", "a:cs"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    latin = rpr.find(qn("a:latin"))
    idx = list(rpr).index(latin)
    if part.get("shadow"):
        shadow = etree.fromstring(
            f'<a:effectLst {NS}><a:outerShdw blurRad="50800" dist="25400" dir="2700000" algn="tl" rotWithShape="0">'
            f'<a:srgbClr val="000000"><a:alpha val="70000"/></a:srgbClr></a:outerShdw></a:effectLst>')
        rpr.insert(idx, shadow)
        idx += 1
    for tag in ("a:cs", "a:ea"):  # 依 schema 順序插在 latin 後面：latin, ea, cs
        el = etree.Element(qn(tag))
        el.set("typeface", font.name)
        rpr.insert(idx + 1, el)


def _textbox(slide, left: int, top: int, width: int, height: int, name: str, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    return tf


def _add_header(slide, song: dict, slide_info: dict, style: dict, w: int, w_in: float) -> None:
    tf = _textbox(slide, Inches(0.5), Inches(style["headerTop"]), w - Inches(1.0), Inches(1.0), "歌名")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    label = slide_info.get("label") or ""
    scale = header_scale(song["title"], label, style, w_in)
    run = p.add_run()
    run.text = song["title"]
    _style_run(run, style["title"], size=style["title"]["size"] * scale)
    if style["showSegment"] and label:
        seg = p.add_run()
        seg.text = "  [" + label + "]"
        _style_run(seg, style["segment"], size=style["segment"]["size"] * scale)


def _add_lyrics(slide, text: str, style: dict, w: int, h: int, w_in: float, h_in: float) -> None:
    left = Inches(style["lyricsLeft"])
    top = Inches(style["lyricsTop"])
    width = max(Inches(1), w - 2 * left)
    height = max(Inches(1), h - top - Inches(0.25))
    tf = _textbox(slide, left, top, width, height, "歌詞")
    size, lines = lyrics_layout(text, style, w_in, h_in)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(style["lyricsAlign"], PP_ALIGN.CENTER)
        p.line_spacing = float(style.get("lineSpacing") or 1.0)
        run = p.add_run()
        run.text = line
        _style_run(run, style["lyrics"], size=size)


def _add_title_page(slide, song: dict, style: dict, w: int, h: int, w_in: float, h_in: float) -> None:
    tf = _textbox(slide, Inches(0.5), 0, w - Inches(1.0), h, "歌名頁", anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = song["title"]
    _style_run(run, style["title"], size=title_page_size(song["title"], style, w_in, h_in))


def _notes(slide, song: dict, info: dict) -> None:
    label = info.get("label") or ""
    if info["kind"] == "title":
        text = f"{song['title']}｜歌名頁"
    else:
        text = f"{song['title']}｜{label or '歌詞'}"
        if info.get("repeat", 1) > 1:
            text += f" ×{info['repeat']}"
        if info["kind"] == "blank":
            text += "（空白頁）"
    if song.get("key"):
        text += f"｜{song['key']}"
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# 主程式
# ---------------------------------------------------------------------------

def build(resolved: dict, images_dir: Path) -> tuple[BytesIO, int, list[str]]:
    """回傳 (pptx 內容, 投影片張數, 注意事項)。"""
    size_key = resolved["service"].get("slideSize", "4:3")
    w_in, h_in = SLIDE_SIZES.get(size_key, SLIDE_SIZES["4:3"])
    prs = Presentation()
    prs.slide_width = Inches(w_in)
    prs.slide_height = Inches(h_in)
    w, h = prs.slide_width, prs.slide_height
    warnings: list[str] = []
    layouts: dict[str, object] = {}
    layout_names: set[str] = set()

    def layout_for(theme: dict, song_title: str):
        image = theme.get("image")
        key = f"{theme.get('id')}|{image['id'] if image else ''}|{theme.get('bgColor')}|{theme.get('overlay')}"
        if key in layouts:
            return layouts[key]
        name = f"背景：{theme.get('name') or '未命名'}"
        n = 2
        while name in layout_names:
            name = f"背景：{theme.get('name') or '未命名'} ({n})"
            n += 1
        layout_names.add(name)
        layout = _new_layout(prs, name)
        image_file = None
        if image:
            try:
                candidate = imagestore.image_path(images_dir, image["id"], image["ext"])
                image_file = candidate if candidate.exists() else None
            except Exception:
                image_file = None
            if image_file is None:
                warnings.append(f"「{song_title}」的背景圖片檔不見了，先用底色代替")
        _set_layout_background(layout, theme, image_file, w_in, h_in)
        _add_layout_overlay(layout, w, h, theme.get("overlay") or 0)
        layouts[key] = layout
        return layout

    count = 0
    for song in resolved["songs"]:
        style, theme = song["style"], song["theme"]
        layout = layout_for(theme, song["title"] or "未命名")
        for w_text in song.get("warnings", []):
            if w_text != "還沒有歌詞":
                warnings.append(f"「{song['title'] or '未命名'}」：{w_text}")
        if not song["slides"]:
            warnings.append(f"「{song['title'] or '未命名'}」沒有歌詞，沒有做投影片")
        for info in song["slides"]:
            slide = prs.slides.add_slide(layout)
            if info["kind"] == "title":
                _add_title_page(slide, song, style, w, h, w_in, h_in)
            elif info["kind"] == "lyrics":
                if style["showHeader"]:
                    _add_header(slide, song, info, style, w, w_in)
                _add_lyrics(slide, info["text"], style, w, h, w_in, h_in)
            _notes(slide, song, info)
            count += 1

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out, count, warnings
