import json
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.dml.color import RGBColor


def hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip("#")
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def create_ppt_from_json(json_path: str):
    with open(json_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    prs = Presentation()

    # 投影比例
    if config.get("slidesize") == "16:9":
        prs.slide_width = Pt(960)
        prs.slide_height = Pt(540)

    for song in config["songs"]:
        bg_image = song["background"]
        font_cfg = song["font"]

        for block in song["lyrics"]:
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            # === 背景圖 ===
            bg = slide.shapes.add_picture(
                bg_image,
                Pt(0),
                Pt(0),
                prs.slide_width,
                prs.slide_height
            )
            slide.shapes._spTree.remove(bg._element)
            slide.shapes._spTree.insert(2, bg._element)

            # === 文字框 ===
            textbox = slide.shapes.add_textbox(
                Pt(50),
                Pt(100),
                prs.slide_width - Pt(100),
                prs.slide_height - Pt(200)
            )

            tf = textbox.text_frame
            tf.clear()
            tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

            p = tf.paragraphs[0]
            p.text = block
            p.alignment = PP_ALIGN.CENTER

            font = p.font
            font.name = font_cfg.get("name", "Arial")
            font.size = Pt(font_cfg.get("size", 56))
            font.bold = True
            font.color.rgb = hex_to_rgb(font_cfg.get("color", "#FFFFFF"))

    output = config.get("output", "敬拜歌單.pptx")
    prs.save(output)
    print(f"✅ PPT 已產生：{output}")


if __name__ == "__main__":
    create_ppt_from_json("songs.json")
