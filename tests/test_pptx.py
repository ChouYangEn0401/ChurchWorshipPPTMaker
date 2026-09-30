import io

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from worship_ppt import pptx_builder, store


def _png(color=(200, 80, 40), size=(1600, 900)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _build(db, paths, service_id):
    resolved = store.resolve_service(db, service_id)
    buf, count, warnings = pptx_builder.build(resolved, paths.images_dir)
    return Presentation(buf), count, warnings, resolved


def test_one_layout_per_background_and_slides_follow_sequence(db, paths, sheet):
    svc = store.import_doc(db, sheet)
    photo = store.add_image(db, paths, _png(), "山景.png")
    store.update_theme(db, photo["id"], {"overlay": 0.4})
    store.apply_look(db, svc["id"], {"themeId": photo["id"]})
    last = svc["items"][3]
    store.update_item(db, last["id"], {"themeId": store.list_themes(db)[-1]["id"]}, None)

    prs, count, warnings, resolved = _build(db, paths, svc["id"])
    assert count == sum(len(s["slides"]) for s in resolved["songs"]) == len(prs.slides) == 42
    layout_names = [s.slide_layout.name for s in prs.slides]
    assert layout_names[0] == "背景：山景"
    assert len(set(layout_names)) == 2  # 兩個背景 → 兩個版面配置，不是每頁一張圖

    layout = prs.slides[0].slide_layout
    bg = layout._element.find(qn("p:cSld")).find(qn("p:bg"))
    blip = bg.find(".//" + qn("a:blip"))
    assert blip is not None
    src_rect = bg.find(".//" + qn("a:srcRect"))
    assert src_rect.get("l") and src_rect.get("r") and not src_rect.get("t")  # 16:9 的圖放進 4:3：左右裁掉，不變形
    overlay = [sp for sp in layout.shapes if sp.name == "暗色遮罩"]
    assert overlay, "暗色遮罩要在版面配置上"
    # 投影片本身沒有背景圖（換背景只要改版面配置）
    assert not any(sh.shape_type == 13 for sh in prs.slides[0].shapes)

    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert notes.startswith("神同在｜主")
    assert any("尾" in w for w in warnings)


def test_text_has_east_asian_font_and_autofit(db, paths, sheet):
    svc = store.import_doc(db, sheet)
    prs, *_ = _build(db, paths, svc["id"])
    lyrics_box = next(sh for sh in prs.slides[0].shapes if sh.name == "歌詞")
    runs = [r for p in lyrics_box.text_frame.paragraphs for r in p.runs]
    rpr = runs[0]._r.find(qn("a:rPr"))
    assert rpr.find(qn("a:ea")).get("typeface") == "Microsoft JhengHei"
    assert rpr.find(qn("a:effectLst")) is not None  # 陰影
    # 「我要跳舞讚美祂 擊鼓彈琴向祂歡呼」差一點點放不下 → 整頁縮小一點，不會把「呼」單獨擠到下一行
    sizes = {r.font.size.pt for r in runs}
    assert len(sizes) == 1 and 37 < sizes.pop() < 44
    assert [p.text for p in lyrics_box.text_frame.paragraphs][2] == "我要跳舞讚美祂 擊鼓彈琴向祂歡呼"


def test_missing_image_falls_back_to_color_with_warning(db, paths, sheet):
    svc = store.import_doc(db, sheet)
    theme = store.add_image(db, paths, _png(), "消失.png")
    store.apply_look(db, svc["id"], {"themeId": theme["id"]})
    for f in paths.images_dir.rglob("*.png"):
        f.unlink()
    prs, count, warnings, _ = _build(db, paths, svc["id"])
    assert count == 42
    assert any("背景圖片檔不見了" in w for w in warnings)


def test_title_slide_and_16_9(db, paths, sheet):
    svc = store.import_doc(db, sheet)
    style = store.create_style(db, {"name": "有歌名頁", "titleSlide": True})
    store.apply_look(db, svc["id"], {"styleId": style["id"]})
    store.update_service(db, svc["id"], {"slideSize": "16:9"}, None)
    prs, count, _, _ = _build(db, paths, svc["id"])
    assert count == 46
    assert round(prs.slide_width / 914400, 2) == 13.33
    first = prs.slides[0]
    assert [sh.name for sh in first.shapes] == ["歌名頁"]
    assert first.shapes[0].text_frame.text == "神同在"


def test_uploaded_webp_is_converted_for_old_powerpoint(db, paths):
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), (10, 20, 30)).save(buf, "WEBP")
    theme = store.add_image(db, paths, buf.getvalue(), "x.webp")
    assert theme["image"]["ext"] == ".jpg"
    big = store.add_image(db, paths, _png(size=(5000, 3000)), "big.png")
    assert big["image"]["width"] == 2560
