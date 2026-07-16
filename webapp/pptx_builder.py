"""
把「已解析好的專案資料」轉成 .pptx。

輸入的 payload 形狀（前端會把每首歌用到的積木直接嵌進去，所以這裡不需要 library）：

    {
      "slideSize": "4:3" | "16:9",
      "songs": [
        {
          "title": "新酒",
          "lyrics": "[Verse]\n...",
          "theme":     { "image": "xxxx.png"|null, "bgColor": "#000000", "overlay": 0.0 },
          "textStyle": {
            "title":   {"font": "...", "size": 60, "color": "#FFD700", "bold": true},
            "segment": {"font": "...", "size": 30, "color": "#FF4500", "bold": false},
            "lyrics":  {"font": "...", "size": 52, "color": "#000000", "bold": false},
            "lyricsLeft": 0.05, "lyricsTop": 1.8,
            "lyricsAlign": "center", "showHeader": true
          }
        }
      ]
    }

一首歌會依 [段落] 標籤被拆成多張投影片，跟原本 main.py 的行為一致。
"""

import os
import re
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn


# 投影片實際尺寸（吋）。前端預覽用同一組數字換算，確保預覽≈成品。
SLIDE_SIZES = {
    "4:3": (10.0, 7.5),
    "16:9": (13.333, 7.5),
}

ALIGN_MAP = {
    "center": PP_ALIGN.CENTER,
    "left": PP_ALIGN.LEFT,
    "right": PP_ALIGN.RIGHT,
}


def hex_to_rgb(hex_color):
    hex_color = (hex_color or "#FFFFFF").lstrip("#")
    if len(hex_color) != 6:
        hex_color = "FFFFFF"
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def parse_lyrics(text):
    """
    解析歌詞成 [(段落名, 該段歌詞), ...]
    依照 [Verse] [Chorus] [Bridge] 等中括號標籤拆段（與原版一致）。
    """
    blocks = []
    current_seg = "Intro"
    current_lines = []
    for line in (text or "").splitlines():
        line = line.rstrip()
        if not line.strip():
            # 保留段落內的空行（可用來排版），但開頭的空行忽略
            if current_lines:
                current_lines.append("")
            continue
        m = re.match(r"\[(.+?)\]\s*$", line.strip())
        if m:
            if current_lines:
                blocks.append((current_seg, "\n".join(current_lines).strip("\n")))
            current_seg = m.group(1)
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines:
        blocks.append((current_seg, "\n".join(current_lines).strip("\n")))
    return blocks


def _set_run_font(run, name, size_pt, hex_color, bold):
    """設定一個 run 的字型/大小/顏色/粗體，並同時套用到 East Asian / Complex Script，
    這樣中文字才會真的用指定字體，而不是掉回 PowerPoint 預設字。"""
    f = run.font
    f.size = Pt(float(size_pt))
    f.bold = bool(bold)
    f.color.rgb = hex_to_rgb(hex_color)
    f.name = name  # 這會建立 <a:latin>

    rPr = run._r.get_or_add_rPr()
    # 移除舊的 ea/cs 再依 schema 順序（latin, ea, cs）插回去
    for tag in ("a:ea", "a:cs"):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    latin = rPr.find(qn("a:latin"))
    idx = (list(rPr).index(latin) + 1) if latin is not None else len(rPr)
    cs = etree.Element(qn("a:cs")); cs.set("typeface", name)
    ea = etree.Element(qn("a:ea")); ea.set("typeface", name)
    rPr.insert(idx, cs)
    rPr.insert(idx, ea)


def _add_overlay(slide, w_emu, h_emu, hex_color, alpha):
    """在背景圖之上鋪一層半透明色塊，讓文字更好讀。alpha 為 0~1。"""
    if not alpha or alpha <= 0:
        return
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, w_emu, h_emu)
    shape.line.fill.background()
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(hex_color)
    # 把 alpha 塞進 <a:srgbClr>
    try:
        spPr = shape.fill._xPr
        srgb = spPr.find(".//" + qn("a:srgbClr"))
        if srgb is not None:
            a = etree.SubElement(srgb, qn("a:alpha"))
            a.set("val", str(int(max(0.0, min(1.0, float(alpha))) * 100000)))
    except Exception:
        pass


def _add_header(slide, title, segment, ts, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = False
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    tstyle = ts.get("title", {})
    sstyle = ts.get("segment", {})

    if title:
        run_t = p.add_run()
        run_t.text = title + "  "
        _set_run_font(run_t, tstyle.get("font", "Arial"), tstyle.get("size", 60),
                      tstyle.get("color", "#FFD700"), tstyle.get("bold", False))
    if segment:
        run_s = p.add_run()
        run_s.text = f"[{segment}]"
        _set_run_font(run_s, sstyle.get("font", "Arial"), sstyle.get("size", 30),
                      sstyle.get("color", "#FF4500"), sstyle.get("bold", False))


def _add_lyrics(slide, content, ts, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()

    lstyle = ts.get("lyrics", {})
    align = ALIGN_MAP.get(ts.get("lyricsAlign", "center"), PP_ALIGN.CENTER)

    lines = (content or "").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _set_run_font(run, lstyle.get("font", "Arial"), lstyle.get("size", 52),
                      lstyle.get("color", "#FFFFFF"), lstyle.get("bold", False))


def build_pptx(payload, uploads_dir):
    """回傳 (BytesIO, warnings)。warnings 是找不到背景圖的歌名清單。"""
    size_key = payload.get("slideSize", "4:3")
    w_in, h_in = SLIDE_SIZES.get(size_key, SLIDE_SIZES["4:3"])

    prs = Presentation()
    prs.slide_width = Inches(w_in)
    prs.slide_height = Inches(h_in)
    blank = prs.slide_layouts[6]

    warnings = []

    for song in payload.get("songs", []):
        title = (song.get("title") or "").strip()
        theme = song.get("theme") or {}
        ts = song.get("textStyle") or {}
        show_header = ts.get("showHeader", True)
        left_in = float(ts.get("lyricsLeft", 0.05) or 0)
        top_in = float(ts.get("lyricsTop", 1.8) or 0)

        blocks = parse_lyrics(song.get("lyrics", ""))
        for seg, lyrics in blocks:
            slide = prs.slides.add_slide(blank)

            # ---- 背景：先底色，再圖片，再遮罩 ----
            bg_color = theme.get("bgColor") or "#000000"
            bg_rect = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg_rect.line.fill.background()
            try:
                bg_rect.shadow.inherit = False
            except Exception:
                pass
            bg_rect.fill.solid()
            bg_rect.fill.fore_color.rgb = hex_to_rgb(bg_color)

            img = theme.get("image")
            if img:
                img_path = os.path.join(uploads_dir, os.path.basename(img))
                if os.path.exists(img_path):
                    try:
                        slide.shapes.add_picture(
                            img_path, 0, 0,
                            width=prs.slide_width, height=prs.slide_height)
                    except Exception:
                        if title not in warnings:
                            warnings.append(title or "(未命名)")
                else:
                    if (title or "(未命名)") not in warnings:
                        warnings.append(title or "(未命名)")

            _add_overlay(slide, prs.slide_width, prs.slide_height,
                         "#000000", theme.get("overlay", 0.0))

            # ---- 標題列 ----
            if show_header:
                _add_header(
                    slide, title, seg, ts,
                    left=Inches(0.5), top=Inches(0.4),
                    width=prs.slide_width - Inches(1), height=Inches(1.0))

            # ---- 歌詞 ----
            _add_lyrics(
                slide, lyrics, ts,
                left=Inches(left_in), top=Inches(top_in),
                width=prs.slide_width - Inches(max(0.0, left_in) * 2),
                height=prs.slide_height - Inches(top_in + 0.2))

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out, warnings
