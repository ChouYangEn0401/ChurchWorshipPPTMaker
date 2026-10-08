import os
import json
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, filedialog, colorchooser, messagebox
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageTk, ImageDraw
import re

# ----------- Utils -----------

COMMON_COLORS = {
    "White": "#FFFFFF",
    "Black": "#000000",
    "Yellow": "#FFD700",
    "Light Yellow": "#FFFACD",
    "Blue": "#1E90FF",
    "Light Blue": "#87CEFA",
    "Red": "#FF4500",
    "Green": "#32CD32",
    "Gray": "#A9A9A9"
}

# 每個文字樣式的預設值（建 UI / 讀檔共用，避免多處各寫一份）
STYLE_ROWS = [
    ("Title", "標題文字設定:  ", 60),
    ("Segment", "段落文字設定:  ", 30),
    ("Lyrics", "歌詞文字設定:  ", 64),
]

# 使用者層級的小設定（記住上次資料夾），放在家目錄
CONFIG_PATH = os.path.join(os.path.expanduser("~"), ".worship_ppt_maker.json")

# 預設投影片尺寸（python-pptx 預設 10 x 7.5 吋），預覽用來換算比例
SLIDE_W_IN = 10.0
SLIDE_H_IN = 7.5


def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def safe_float(var, default=0.0):
    """從 tk 變數安全取出非負浮點數，讀不到或負數就回傳預設值。"""
    try:
        v = float(var.get())
        return v if v >= 0 else default
    except (ValueError, tk.TclError):
        return default


def parse_lyrics(text):
    """
    解析歌詞成 [(SEG, lyrics), ...] 格式
    依照 [Verse], [Chorus], [Bridge], [Pre-Chorus] 等標籤拆段
    """
    blocks = []
    current_seg = "Intro"
    current_lines = []

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.match(r'\[(.+?)\]', line)
        if match:
            # 新段落
            if current_lines:
                blocks.append((current_seg, "\n".join(current_lines)))
            current_seg = match.group(1)
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines:
        blocks.append((current_seg, "\n".join(current_lines)))
    return blocks


# ----------- Closable Notebook（可點擊 ✖ 關閉分頁）-----------

class ClosableNotebook(ttk.Notebook):
    """在每個分頁右側畫一個真的可以點擊的 ✖，點下去就關閉該分頁。"""

    _style_ready = False
    _imgs = None  # 保存 icon 參照，避免被回收

    def __init__(self, master, **kw):
        if not ClosableNotebook._style_ready:
            self._init_style()
            ClosableNotebook._style_ready = True
        kw["style"] = "ClosableNotebook"
        super().__init__(master, **kw)
        self._pressed_index = None
        self.bind("<ButtonPress-1>", self._on_press, add=True)
        self.bind("<ButtonRelease-1>", self._on_release, add=True)

    @staticmethod
    def _make_icon(color):
        img = Image.new("RGBA", (14, 14), (0, 0, 0, 0))
        d = ImageDraw.Draw(img)
        d.line((3, 3, 10, 10), fill=color, width=2)
        d.line((10, 3, 3, 10), fill=color, width=2)
        return ImageTk.PhotoImage(img)

    def _init_style(self):
        normal = self._make_icon((120, 120, 120, 255))
        active = self._make_icon((210, 60, 60, 255))
        pressed = self._make_icon((150, 20, 20, 255))
        ClosableNotebook._imgs = (normal, active, pressed)

        style = ttk.Style()
        style.element_create(
            "close", "image", str(normal),
            ("active", "pressed", "!disabled", str(pressed)),
            ("active", "!disabled", str(active)),
            border=6, sticky="",
        )
        style.layout("ClosableNotebook", [("ClosableNotebook.client", {"sticky": "nswe"})])
        style.layout("ClosableNotebook.Tab", [
            ("ClosableNotebook.tab", {"sticky": "nswe", "children": [
                ("ClosableNotebook.padding", {"side": "top", "sticky": "nswe", "children": [
                    ("ClosableNotebook.focus", {"side": "top", "sticky": "nswe", "children": [
                        ("ClosableNotebook.label", {"side": "left", "sticky": ""}),
                        ("ClosableNotebook.close", {"side": "left", "sticky": ""}),
                    ]})
                ]})
            ]})
        ])

    def _on_press(self, event):
        element = self.identify(event.x, event.y)
        if "close" in element:
            self._pressed_index = self.index("@%d,%d" % (event.x, event.y))
            self.state(["pressed"])
            return "break"

    def _on_release(self, event):
        if not self.instate(["pressed"]):
            return
        element = self.identify(event.x, event.y)
        self.state(["!pressed"])
        if "close" in element:
            index = self.index("@%d,%d" % (event.x, event.y))
            if index == self._pressed_index:
                self.forget(index)
                self.event_generate("<<NotebookTabClosed>>")
        self._pressed_index = None


# ----------- Page Tab -----------

class SongTab(ttk.Frame):
    def __init__(self, master, clone_from=None):
        super().__init__(master)
        self.bg_image = tk.StringVar()
        self.lyrics_left = tk.DoubleVar(value=0.05)
        self.lyrics_top = tk.DoubleVar(value=1.8)

        self._build_ui()
        if clone_from:
            self.clone_settings(clone_from)

    # ---- 可捲動的內容容器：控制項多、視窗變矮時不會被切掉 ----
    def _build_scroll_body(self):
        canvas = tk.Canvas(self, highlightthickness=0)
        vsb = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        body = ttk.Frame(canvas)
        win = canvas.create_window((0, 0), window=body, anchor="nw")

        body.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(win, width=e.width))

        self.scroll_canvas = canvas
        return body

    def _build_ui(self):
        body = self._build_scroll_body()

        # ----- 歌詞輸入 -----
        lyrics_frame = ttk.LabelFrame(body, text="歌曲資訊")
        lyrics_frame.pack(fill="both", expand=True, padx=10, pady=5)

        ttk.Label(lyrics_frame, text="歌曲名稱: ").pack(anchor="w")
        self.title_text = tk.Text(lyrics_frame, height=1)
        self.title_text.pack(fill="x", pady=2)

        ttk.Label(lyrics_frame, text="所有歌詞:\n\t[分段]歌詞 -> 製成一頁PPT").pack(anchor="w", pady=(5, 0))
        self.lyrics_text = tk.Text(lyrics_frame, height=15)
        self.lyrics_text.pack(fill="both", expand=True, pady=2)

        # ----- Style Controls -----
        style_frame = ttk.LabelFrame(body, text="PPT字幕顏色")
        style_frame.pack(fill="x", padx=10, pady=5)

        self.styles = {}
        for key, name, default_size in STYLE_ROWS:
            self._style_row(style_frame, key, name, default_size)

        # ----- Lyric Position Controls -----
        pos_frame = ttk.LabelFrame(body, text="歌詞位置調整 (吋)")
        pos_frame.pack(fill="x", padx=10, pady=5)

        vcmd = (self.register(self._validate_pos), "%P")
        ttk.Label(pos_frame, text="距離左邊:  ").pack(side="left")
        ttk.Entry(pos_frame, textvariable=self.lyrics_left, width=6,
                  validate="key", validatecommand=vcmd).pack(side="left", padx=5)
        ttk.Label(pos_frame, text="距離上邊:  ").pack(side="left", padx=(10, 0))
        ttk.Entry(pos_frame, textvariable=self.lyrics_top, width=6,
                  validate="key", validatecommand=vcmd).pack(side="left", padx=5)
        ttk.Label(pos_frame, text="(只能填 0 或正數)").pack(side="left", padx=(10, 0))

        # ----- Background -----
        bg_frame = ttk.LabelFrame(body, text="背景設定")
        bg_frame.pack(fill="x", padx=10, pady=5)
        ttk.Button(bg_frame, text="選擇背景圖片", command=self.select_bg).pack(side="left")
        ttk.Label(bg_frame, textvariable=self.bg_image).pack(side="left", padx=10)

    @staticmethod
    def _validate_pos(proposed):
        # 允許逐字輸入（空字串 / 單一小數點），但擋掉負數與非數字
        if proposed in ("", "."):
            return True
        try:
            return float(proposed) >= 0
        except ValueError:
            return False

    def _style_row(self, parent, key, name, default_size):
        frame = ttk.Frame(parent)
        frame.pack(fill="x", pady=2)
        ttk.Label(frame, text=name, width=15).pack(side="left")

        font_var = tk.StringVar(value="Arial")
        size_var = tk.IntVar(value=default_size)
        color_var = tk.StringVar(value="#FFFFFF")

        ttk.Entry(frame, textvariable=font_var, width=10).pack(side="left", padx=3)
        ttk.Entry(frame, textvariable=size_var, width=5).pack(side="left", padx=3)

        btn = tk.Button(frame, width=3, relief="solid", bd=1,
                        bg=color_var.get(), activebackground=color_var.get())
        btn.pack(side="left", padx=3)
        btn.bind("<Button-1>", lambda e, v=color_var, b=btn: self.pick_common_color(v, b))
        btn.bind("<Button-3>", lambda e, v=color_var, b=btn: self.pick_color_picker(v, b))

        self.styles[key] = {"font": font_var, "size": size_var, "color": color_var, "_btn": btn}

    def pick_common_color(self, var, btn):
        menu = tk.Menu(self, tearoff=0)
        for name, hexcode in COMMON_COLORS.items():
            menu.add_command(label=name, background=hexcode,
                             command=lambda h=hexcode: self._set_color(var, btn, h))
        menu.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def pick_color_picker(self, var, btn):
        c = colorchooser.askcolor()[1]
        if c:
            self._set_color(var, btn, c)

    def _set_color(self, var, btn, color):
        var.set(color)
        # 同時設 activebackground，避免按下時露出預設灰色（原本的「秀逗」）
        btn.config(bg=color, activebackground=color)

    def select_bg(self):
        path = filedialog.askopenfilename(
            initialdir=App.get_last_dir(),
            filetypes=[("Image", "*.jpg *.png *.jpeg")])
        if path:
            self.bg_image.set(path)
            App.set_last_dir(path)

    def clone_settings(self, other):
        # 只複製「外觀設定」(字型/顏色/位置/背景)，不複製歌名與歌詞內容
        self.bg_image.set(other.bg_image.get())
        self.lyrics_left.set(other.lyrics_left.get())
        self.lyrics_top.set(other.lyrics_top.get())
        for k in self.styles:
            self.styles[k]["font"].set(other.styles[k]["font"].get())
            self.styles[k]["size"].set(other.styles[k]["size"].get())
            self.styles[k]["color"].set(other.styles[k]["color"].get())


# ----------- Main App -----------

class WorshipPPTApp(tk.Tk):
    _last_dir = ""

    def __init__(self):
        super().__init__()
        self.title("教會 敬拜歌曲PPT 快速生成器")
        self.geometry("950x700")

        WorshipPPTApp._last_dir = self._load_last_dir()

        self._menu()

        # Notebook（可點擊 ✖ 關閉分頁）
        self.notebook = ClosableNotebook(self)
        self.notebook.pack(fill="both", expand=True)

        # 滑鼠滾輪捲動目前分頁（在歌詞框上則交給歌詞框自己捲）
        self.bind_all("<MouseWheel>", self._on_mousewheel)

        self.add_tab()

    # ---------- 使用者設定：記住上次資料夾 ----------
    @classmethod
    def get_last_dir(cls):
        return cls._last_dir or ""

    @classmethod
    def set_last_dir(cls, path):
        cls._last_dir = os.path.dirname(path)
        try:
            with open(CONFIG_PATH, "w", encoding="utf-8") as f:
                json.dump({"last_dir": cls._last_dir}, f)
        except OSError:
            pass

    @staticmethod
    def _load_last_dir():
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                return json.load(f).get("last_dir", "")
        except (OSError, ValueError):
            return ""

    def _menu(self):
        menu_frame = ttk.Frame(self)
        menu_frame.pack(fill="x")
        ttk.Button(menu_frame, text="+ 新增歌曲", command=self.add_tab).pack(side="left", padx=5)
        ttk.Button(menu_frame, text="生成 PPT", command=self.generate_ppt).pack(side="left", padx=5)
        ttk.Button(menu_frame, text="預覽", command=self.open_preview).pack(side="left", padx=5)

        ttk.Button(menu_frame, text="儲存專案", command=self.save_json).pack(side="left", padx=20)
        ttk.Button(menu_frame, text="讀取專案 (將覆蓋掉現在沒有存檔的一切)",
                   command=self.load_json).pack(side="left", padx=5)

    def _on_mousewheel(self, event):
        w = self.winfo_containing(event.x_root, event.y_root)
        if isinstance(w, tk.Text):
            return  # 在文字框上時，讓文字框自己捲動內容
        try:
            tab = self.notebook.nametowidget(self.notebook.select())
        except (tk.TclError, KeyError):
            return
        canvas = getattr(tab, "scroll_canvas", None)
        if canvas is not None:
            canvas.yview_scroll(int(-event.delta / 120), "units")

    # ---------- 分頁標題即時顯示歌名 ----------
    def _bind_tab_title(self, tab):
        tab.title_text.bind("<KeyRelease>", lambda e, t=tab: self._refresh_tab_text(t))

    def _refresh_tab_text(self, tab):
        title = tab.title_text.get("1.0", "end").strip()
        try:
            self.notebook.tab(tab, text=(title or "未命名"))
        except tk.TclError:
            pass

    def save_json(self):
        data = {"songs": []}

        for tab_id in self.notebook.tabs():
            tab = self.notebook.nametowidget(tab_id)

            song = {
                "title": tab.title_text.get("1.0", "end").strip(),
                "lyrics": tab.lyrics_text.get("1.0", "end").strip(),
                "background": tab.bg_image.get(),
                "lyrics_left": safe_float(tab.lyrics_left, 0.05),
                "lyrics_top": safe_float(tab.lyrics_top, 1.8),
                "styles": {}
            }

            for key, style in tab.styles.items():
                song["styles"][key] = {
                    "font": style["font"].get(),
                    "size": style["size"].get(),
                    "color": style["color"].get(),
                }

            data["songs"].append(song)

        path = filedialog.asksaveasfilename(
            initialdir=self.get_last_dir(),
            defaultextension=".json",
            filetypes=[("JSON File", "*.json")]
        )

        if path:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            self.set_last_dir(path)
            messagebox.showinfo("Saved", "Project saved successfully.")

    def load_json(self):
        path = filedialog.askopenfilename(
            initialdir=self.get_last_dir(),
            filetypes=[("JSON File", "*.json")]
        )

        if not path:
            return

        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        self.set_last_dir(path)

        # ---------- 清空現有 tabs ----------
        for tab_id in self.notebook.tabs():
            self.notebook.forget(tab_id)

        # ---------- 重建 tabs ----------
        for song in data.get("songs", []):
            tab = SongTab(self.notebook)
            title = song.get("title", "")
            self.notebook.add(tab, text=(title or "未命名"))

            # Title / Lyrics
            tab.title_text.insert("1.0", title)
            tab.lyrics_text.insert("1.0", song.get("lyrics", ""))

            # Background
            tab.bg_image.set(song.get("background", ""))
            tab.lyrics_left.set(song.get("lyrics_left", 0.05))
            tab.lyrics_top.set(song.get("lyrics_top", 1.8))

            # Styles
            for key, style in song.get("styles", {}).items():
                if key in tab.styles:
                    tab.styles[key]["font"].set(style.get("font", "Arial"))
                    tab.styles[key]["size"].set(style.get("size", 40))
                    color = style.get("color", "#FFFFFF")
                    tab.styles[key]["color"].set(color)

            self._bind_tab_title(tab)
            self._sync_color_buttons(tab)

        if self.notebook.tabs():
            self.notebook.select(self.notebook.tabs()[0])

        messagebox.showinfo("Loaded", "Project loaded successfully.")

    @staticmethod
    def _sync_color_buttons(tab):
        # 讀檔後把每個顏色按鈕的底色更新成實際顏色
        for key, style in tab.styles.items():
            btn = style.get("_btn")
            if btn is not None:
                c = style["color"].get()
                btn.config(bg=c, activebackground=c)

    def add_tab(self):
        clone = None
        if self.notebook.tabs():
            clone = self.notebook.nametowidget(self.notebook.select())

        tab = SongTab(self.notebook, clone_from=clone)
        self.notebook.add(tab, text="未命名")
        self.notebook.select(tab)
        self._bind_tab_title(tab)
        self._sync_color_buttons(tab)

    # ---------- 收集所有將輸出的投影片（生成 / 預覽共用）----------
    def _iter_slides(self):
        slides = []
        for tab_id in self.notebook.tabs():
            tab = self.notebook.nametowidget(tab_id)
            title = tab.title_text.get("1.0", "end").strip()
            blocks = parse_lyrics(tab.lyrics_text.get("1.0", "end"))
            for seg, lyrics in blocks:
                slides.append({"tab": tab, "title": title, "seg": seg, "lyrics": lyrics})
        return slides

    def generate_ppt(self):
        slides = self._iter_slides()
        if not slides:
            messagebox.showwarning("沒有內容", "目前沒有任何歌詞可以生成。")
            return

        out_path = filedialog.asksaveasfilename(
            initialdir=self.get_last_dir(),
            defaultextension=".pptx",
            filetypes=[("PowerPoint", "*.pptx")]
        )
        if not out_path:
            return

        # ---------- 進度視窗 ----------
        prog = tk.Toplevel(self)
        prog.title("生成中...")
        prog.transient(self)
        prog.resizable(False, False)
        ttk.Label(prog, text="正在生成 PPT，請稍候...").pack(padx=20, pady=(15, 5))
        bar = ttk.Progressbar(prog, maximum=len(slides), length=320)
        bar.pack(padx=20, pady=(0, 15))
        prog.update()

        prs = Presentation()
        blank = prs.slide_layouts[6]  # Blank slide
        warned_bg_missing = set()

        for i, s in enumerate(slides):
            tab = s["tab"]
            title = s["title"]
            slide = prs.slides.add_slide(blank)

            # ---------- Background ----------
            bg_path = tab.bg_image.get()
            song_name = title or f"Song {i + 1}"

            if bg_path:
                try:
                    slide.shapes.add_picture(
                        bg_path, 0, 0,
                        width=prs.slide_width, height=prs.slide_height)
                except (FileNotFoundError, OSError):
                    if song_name not in warned_bg_missing:
                        warned_bg_missing.add(song_name)
                        messagebox.showwarning(
                            "背景圖片不存在",
                            f"找不到背景圖片，該歌曲將不使用背景圖。\n\n"
                            f"歌曲：{song_name}\n檔案：\n{bg_path}")

            # ---------- HEADER : Title + Segment (同一行) ----------
            self.add_header(
                slide,
                title=title,
                segment=s["seg"],
                title_style=tab.styles["Title"],
                segment_style=tab.styles["Segment"],
                left=Inches(0.5),
                top=Inches(0.4),
                width=prs.slide_width - Inches(1),
                height=Inches(1.0)
            )

            # ---------- LYRICS : 滿版 ----------
            left = safe_float(tab.lyrics_left, 0.05)
            top = safe_float(tab.lyrics_top, 1.8)
            self.add_textbox(
                slide,
                s["lyrics"],
                tab.styles["Lyrics"],
                left=Inches(left),
                top=Inches(top),
                width=prs.slide_width - Inches(left * 2),
                height=prs.slide_height - Inches(2.1)
            )

            bar["value"] = i + 1
            prog.update()

        prs.save(out_path)
        prog.destroy()
        messagebox.showinfo("Done", f"PPT generated:\n{out_path}")

    # ---------- 預覽視窗：在畫布上模擬投影片外觀 ----------
    def open_preview(self):
        slides = self._iter_slides()
        if not slides:
            messagebox.showinfo("預覽", "目前沒有可預覽的內容。")
            return

        win = tk.Toplevel(self)
        win.title("預覽（重新開啟以更新內容）")
        win.resizable(False, False)

        CW, CH = 720, 540  # 4:3，對齊 10 x 7.5 吋
        ppi = CW / SLIDE_W_IN

        canvas = tk.Canvas(win, width=CW, height=CH, bg="black", highlightthickness=0)
        canvas.pack(padx=10, pady=10)

        nav = ttk.Frame(win)
        nav.pack(pady=(0, 10))
        counter = ttk.Label(nav, text="")

        state = {"i": 0, "bg_cache": {}, "fonts": []}

        def font_px(style_dict, min_px=6):
            size_pt = style_dict["size"].get()
            f = tkfont.Font(family=style_dict["font"].get(),
                            size=-max(min_px, int(size_pt * ppi / 72)))
            state["fonts"].append(f)  # 保存參照，避免重繪時字型被回收
            return f

        def render():
            canvas.delete("all")
            state["fonts"].clear()
            s = slides[state["i"]]
            tab = s["tab"]

            # 背景
            bg = tab.bg_image.get()
            if bg:
                try:
                    if bg not in state["bg_cache"]:
                        im = Image.open(bg).convert("RGB").resize((CW, CH))
                        state["bg_cache"][bg] = ImageTk.PhotoImage(im)
                    canvas.create_image(0, 0, anchor="nw", image=state["bg_cache"][bg])
                except (FileNotFoundError, OSError):
                    canvas.create_text(CW / 2, 16, text="（背景圖遺失）", fill="#FF5555")

            # 標題 + 段落（同一行置中，各自套用自己的字型顏色）
            tstyle = tab.styles["Title"]
            sstyle = tab.styles["Segment"]
            tfont = font_px(tstyle)
            sfont = font_px(sstyle)
            t_txt = (s["title"] + "  ") if s["title"] else ""
            s_txt = f"[{s['seg']}]"
            tw = tfont.measure(t_txt)
            sw = sfont.measure(s_txt)
            y = 0.4 * ppi
            x0 = CW / 2 - (tw + sw) / 2
            if t_txt:
                canvas.create_text(x0, y, anchor="nw", text=t_txt,
                                   font=tfont, fill=tstyle["color"].get())
            canvas.create_text(x0 + tw, y, anchor="nw", text=s_txt,
                               font=sfont, fill=sstyle["color"].get())

            # 歌詞（置中、依設定換行）
            lstyle = tab.styles["Lyrics"]
            lfont = font_px(lstyle)
            left = safe_float(tab.lyrics_left, 0.05)
            top = safe_float(tab.lyrics_top, 1.8)
            wpx = max(1, (SLIDE_W_IN - 2 * left) * ppi)
            canvas.create_text(CW / 2, top * ppi, anchor="n", text=s["lyrics"],
                               font=lfont, fill=lstyle["color"].get(),
                               width=wpx, justify="center")

            counter.config(text=f"{state['i'] + 1} / {len(slides)}")

        def go(step):
            state["i"] = (state["i"] + step) % len(slides)
            render()

        ttk.Button(nav, text="◀ 上一頁", command=lambda: go(-1)).pack(side="left", padx=5)
        counter.pack(side="left", padx=10)
        ttk.Button(nav, text="下一頁 ▶", command=lambda: go(1)).pack(side="left", padx=5)
        win.bind("<Left>", lambda e: go(-1))
        win.bind("<Right>", lambda e: go(1))

        render()

    def add_header(self, slide, title, segment, title_style, segment_style, left, top, width, height):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = False
        tf.clear()

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER

        # ----- TITLE -----
        run_title = p.add_run()
        run_title.text = title + "  "
        run_title.font.name = title_style["font"].get()
        run_title.font.size = Pt(title_style["size"].get())
        run_title.font.color.rgb = hex_to_rgb(title_style["color"].get())

        # ----- SEGMENT -----
        run_seg = p.add_run()
        run_seg.text = f"[{segment}]"
        run_seg.font.name = segment_style["font"].get()
        run_seg.font.size = Pt(segment_style["size"].get())
        run_seg.font.color.rgb = hex_to_rgb(segment_style["color"].get())

    def add_textbox(self, slide, content, style, left, top, width, height):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()

        # 整個文字都是同一樣式
        p = tf.paragraphs[0]
        p.text = content.strip()
        p.alignment = PP_ALIGN.CENTER

        for run in p.runs:
            run.font.name = style["font"].get()
            run.font.size = Pt(style["size"].get())
            run.font.color.rgb = hex_to_rgb(style["color"].get())

        # 如果只有一個段落，就確保整個段落套用顏色
        if len(tf.paragraphs) == 1:
            for paragraph in tf.paragraphs:
                for r in paragraph.runs:
                    r.font.color.rgb = hex_to_rgb(style["color"].get())


# 讓 SongTab 內的靜態方法能參照到 App 的 last_dir helper
App = WorshipPPTApp


# ----------- Run App -----------

if __name__ == "__main__":
    WorshipPPTApp().mainloop()
